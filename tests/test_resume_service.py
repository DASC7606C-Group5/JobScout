"""Real PDF/Word extraction, malformed documents and resource limits."""

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from jobscout.services.resume_service import (
    MAX_RESUME_BYTES,
    MAX_RESUME_TEXT_LENGTH,
    ResumeParseError,
    parse_resume,
)


def make_pdf(*pages: str, encrypted: bool = False) -> bytes:
    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=595, height=842)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 50 750 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = stream
    if encrypted:
        writer.encrypt("test-password")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_preserves_page_order_and_original_name() -> None:
    parsed = parse_resume("resume.PDF", make_pdf("Skills: Python, SQL", "Projects: JobScout"))
    assert parsed.name == "resume.PDF"
    assert parsed.text == "Skills: Python, SQL\nProjects: JobScout"


def test_txt_handles_bom_line_endings_and_browser_paths() -> None:
    parsed = parse_resume("C:\\fakepath\\简历.TXT", "\ufeff 技能\r\nPython \r".encode())
    assert parsed.name == "简历.TXT"
    assert parsed.text == "技能\nPython"


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("教育背景")
    document.add_paragraph("计算机科学学士")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "技能"
    table.cell(0, 1).text = "Python, SQL"
    table.cell(1, 0).text = "项目经历"
    table.cell(1, 1).text = "求职推荐系统"
    document.add_paragraph("实习经历")
    document.add_paragraph("软件工程实习生")
    document.sections[0].header.paragraphs[0].text = "张同学"
    document.sections[0].footer.paragraphs[0].text = "联系邮箱：example@example.test"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_preserves_chinese_paragraphs_tables_and_headers() -> None:
    parsed = parse_resume("简历.DOCX", make_docx())
    assert parsed.name == "简历.DOCX"
    assert parsed.text.splitlines() == [
        "教育背景",
        "计算机科学学士",
        "技能",
        "Python, SQL",
        "项目经历",
        "求职推荐系统",
        "实习经历",
        "软件工程实习生",
        "张同学",
        "联系邮箱：example@example.test",
    ]


def test_docx_reads_nested_tables_and_merged_cells_once() -> None:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.text = "Skills"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Python, SQL"
    output = BytesIO()
    document.save(output)
    parsed = parse_resume("resume.docx", output.getvalue())
    assert parsed.text.count("Skills") == 1
    assert "Python, SQL" in parsed.text


def test_docx_rejects_wrong_zip_and_excessive_expansion() -> None:
    for entry_name, data, code in [
        ("not-a-document.txt", b"example", "invalid_file"),
        ("word/document.xml", b"x" * (20 * 1024 * 1024 + 1), "document_too_large"),
    ]:
        output = BytesIO()
        with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
            archive.writestr(entry_name, data)
        with pytest.raises(ResumeParseError) as raised:
            parse_resume("resume.docx", output.getvalue())
        assert raised.value.code == code


@pytest.mark.parametrize(
    ("name", "content", "code", "status"),
    [
        ("resume.doc", b"legacy", "unsupported_format", 415),
        ("resume.docx", b"word", "invalid_file", 422),
        ("resume.rtf", b"unsupported", "unsupported_format", 415),
        (None, b"unknown", "unsupported_format", 415),
        ("resume.txt", b"", "empty_file", 422),
        ("resume.txt", b"  \n", "empty_text", 422),
        ("resume.txt", b"\xff", "invalid_encoding", 422),
        ("resume.txt", b"bad\x00text", "invalid_text", 422),
        ("resume.txt", b"x" * (MAX_RESUME_BYTES + 1), "file_too_large", 413),
        ("resume.txt", b"x" * (MAX_RESUME_TEXT_LENGTH + 1), "text_too_long", 422),
        ("resume.pdf", b"plain text with wrong extension", "invalid_file", 422),
        ("resume.pdf", b"%PDF-1.7\ninvalid", "invalid_file", 422),
    ],
    ids=[
        "doc",
        "docx",
        "rtf",
        "missing-name",
        "empty-file",
        "empty-text",
        "invalid-encoding",
        "binary-text",
        "oversized-file",
        "oversized-text",
        "fake-pdf",
        "corrupt-pdf",
    ],
)
def test_invalid_resume_fails_with_a_useful_error(
    name: str | None, content: bytes, code: str, status: int
) -> None:
    with pytest.raises(ResumeParseError) as raised:
        parse_resume(name, content)
    assert raised.value.code == code
    assert raised.value.status_code == status


def test_pdf_without_text_prompts_for_ocr() -> None:
    with pytest.raises(ResumeParseError, match="OCR") as raised:
        parse_resume("scanned.pdf", make_pdf(""))
    assert raised.value.code == "no_extractable_text"


def test_encrypted_pdf_prompts_for_password_removal() -> None:
    with pytest.raises(ResumeParseError, match="password") as raised:
        parse_resume("encrypted.pdf", make_pdf("Skills: Python", encrypted=True))
    assert raised.value.code == "encrypted_file"


def test_pdf_page_limit() -> None:
    with pytest.raises(ResumeParseError) as raised:
        parse_resume("long.pdf", make_pdf(*(["Skills: Python"] * 51)))
    assert raised.value.code == "document_too_large"
